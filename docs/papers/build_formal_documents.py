"""Build editable, two-column working papers from the two Formal Markdown sources.

Dependencies: python -m pip install --target .paper_tools pypandoc_binary python-docx
Run from the project root: python docs/papers/build_formal_documents.py
"""
from pathlib import Path
import re
import sys
import json
import argparse
import html
from copy import deepcopy

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.paper_tools'))
import pypandoc
from docx import Document

OUT = ROOT / 'docs/papers'
AUTHORS = json.loads((OUT / 'authors.json').read_text(encoding='utf-8'))


def plain_html(value):
    return html.unescape(re.sub(r'<[^>]+>', '', value)).strip()


def extract_figures(source, english):
    figures = {}

    def replace(match):
        block = match.group()
        caption = plain_html(re.search(r'<figcaption>(.*?)</figcaption>', block, re.S).group(1))
        number = int(re.search(r'(?:Fig\.|图)\s*(\d+)', caption).group(1))
        panels = []
        for cell in re.findall(r'<td[^>]*>(.*?)</td>', block, re.S):
            img = re.search(r'<img\s+src="([^"]+)"[^>]*>', cell)
            if img:
                panels.append((img.group(1), plain_html(cell[img.end():])))
        figures[number] = {'caption': caption, 'panels': panels}
        return f'\nFIGUREGROUP{number}\n'

    source = re.sub(r'<figure\b[^>]*>.*?</figure>', replace, source, flags=re.S)
    # Separate consecutive full-width blocks with the related body text.
    # Word cannot float these blocks, so grouping them leaves large blank areas.
    if 'FIGUREGROUP1' in source:
        source = source.replace('FIGUREGROUP1', '')
        source = re.sub(r'(?=^### 3\.4 )', 'FIGUREGROUP1\n\n',
                        source, count=1, flags=re.M)
    if 'FIGUREGROUP4' in source:
        source = source.replace('FIGUREGROUP4', '')
        source = re.sub(r'(?=^### 6\.2 )', 'FIGUREGROUP4\n\n',
                        source, count=1, flags=re.M)
    # The wide overview already identifies the microscope. Remove the
    # redundant close-up so the environment figure fits a compact 2x2 panel.
    if 1 in figures:
        figures[1]['panels'].pop(1)
        figures[1]['panels'] = [(path, re.sub(r'^\([a-z]\)', '(' + chr(97+i) + ')', label))
                                 for i, (path, label) in enumerate(figures[1]['panels'])]
    # Integrate the speech screenshot into the interaction figure instead of
    # leaving an unnumbered, full-size screenshot in the architecture section.
    source = re.sub(r'\*\*(?:Runtime speech-input evidence|语音输入运行时证据).*?\n\n!\[.*?\]\(<.*?>\)', '', source, flags=re.S)
    if 4 in figures:
        figures[4]['panels'][-1] = (
            'Pic/3.5 语音转文字.png',
            '(d) Speech-input interface.' if english else '(d) 语音输入界面。')
        figures[4]['caption'] = (
            'Fig. 4. Representative assistant interfaces: (a) component description and experiment entry; '
            '(b) contextual explanation; (c) spatial target guidance; and (d) speech input. '
            'These views illustrate interface states, not one uninterrupted session. '
            'The associated NA experiment feedback is shown in Fig. 2(c).'
            if english else
            '图 4. 助手的代表性交互界面：(a) 部件说明与实验入口；(b) 上下文解释；'
            '(c) 空间目标引导；(d) 语音输入。各图表示界面状态，不代表一次连续会话。关联 NA 实验的反馈见图 2(c)。')
    return source, figures


def clean(source, english):
    source, figures = extract_figures(source, english)
    notes = '\n'.join(line for line in source.splitlines() if line.startswith('>'))
    lines = []
    for line in source.splitlines():
        if line.startswith('>'):
            caption = re.match(r'> \*\*(?:Proposed caption:|拟图注：)\*\* (.*)', line)
            if caption:
                lines.extend(['', caption.group(1), ''])
            else:
                figure = re.search(r'(?:FIGURE |插图 )([1-4])(?:\s|｜)', line)
                if figure:
                    n = figure.group(1)
                    lines.extend(['', f'FIGURE_SLOT_{n}', ''])
            continue
        lines.append(line)
    text = '\n'.join(lines)
    if english:
        text = text.replace('This draft retains completion counts and questionnaire means as descriptive results.',
                            'Completion counts and questionnaire means are reported descriptively.')
        text = text.replace('Identify any additional generative tools used for text, figures, or code and confirm the final disclosure against actual use.', '')
    else:
        text = re.sub(r'请?根据实际使用情况.*?。', '', text)
        text = text.replace('若还使用其他生成式工具制作正文、图像或代码，应补充工具和用途；最终声明应按实际使用情况确认。', '')
    # IEEE uses an inline abstract lead-in rather than a section heading.
    text = re.sub(r'## (Abstract|摘要)\s*\n\s*\n', lambda m: '**' + m.group(1) + '—**', text)
    text = text.replace('**Keywords:**', '**Index Terms—**').replace('**关键词：**', '**关键词—**')
    return text, notes, figures


def preserve_display_equations(target, previous):
    """Keep the previously repaired Word equation XML when its math is unchanged."""
    if not previous.exists():
        return
    old, new = Document(previous), Document(target)
    old_blocks = [p for p in old.paragraphs if p._p.xpath('.//m:oMathPara')]
    new_blocks = [p for p in new.paragraphs if p._p.xpath('.//m:oMathPara')]
    if len(old_blocks) != len(new_blocks):
        raise ValueError('Display-equation count changed; review before exporting.')
    for before, after in zip(old_blocks, new_blocks):
        # Pandoc can introduce empty alignment slots, but the math tokens must agree.
        math_text = lambda p: ''.join(p._p.xpath('.//m:t/text()')).replace('\u2001', '').replace(' ', '')
        if math_text(before) != math_text(after):
            raise ValueError('Equation content differs; do not restore an outdated equation.')
        after._p.getparent().replace(after._p, deepcopy(before._p))
    new.save(target)


def anonymize(text, english):
    """Remove identity-bearing front matter and institutional wording for review."""
    if english:
        front = r'^\*\*Authors:\*\*.*\n\*\*Affiliations:\*\*.*\n\*\*Corresponding Authors:\*\*.*$'
        text, count = re.subn(front, 'Anonymous authors', text, count=1, flags=re.M)
        text = text.replace(
            'In accordance with institutional research ethics guidelines at UESTC,',
            "In accordance with the research ethics guidelines of the authors' institution,")
    else:
        front = r'^\*\*作者：\*\*.*\n\*\*单位：\*\*.*\n\*\*通讯作者：\*\*.*$'
        text, count = re.subn(front, '匿名作者', text, count=1, flags=re.M)
        text = text.replace(
            '根据电子科技大学（UESTC）机构研究伦理指南，',
            '根据作者所在机构的研究伦理指南，')
    if count != 1:
        raise ValueError('Could not locate the author block for anonymous export.')
    return text


def main():
    from ieee_layout import style_document
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--anonymous', action='store_true', help='Generate a review copy without author identities.')
    parser.add_argument('--output-root', type=Path, default=OUT, help='Root for formal/anonymous export folders.')
    args = parser.parse_args()
    checklist = ['# 论文排版与待补材料\n',
        '已生成中英文双栏 Word 排版稿，公式为可编辑 Word 数学对象；图 3 已同步至 Formal.md。',
        '图 1、图 2 和图 4 为运行时截图；图 3 已由 draw.io 绘制并嵌入 SVG（附 PNG 兼容图）。伦理声明沿用作者提供的文本，研究记录仍由作者核对；最终页数另见 FINAL_REVIEW.md。',
        '第 6.6 节及表 5 保留在正文，包含 78 次后端请求与 75/78 条人工认可结果；澄清修订后的真实模型重测尚未报告。\n']
    for name, english in [('AIxVR2027_VRMicroscope_Draft_Formal', False), ('AIxVR2027_VRMicroscope_Paper_EN_Formal', True)]:
        source = (OUT / (name+'.md')).read_text(encoding='utf-8-sig')
        content, notes, figures = clean(source, english)
        if args.anonymous:
            content = anonymize(content, english)
        for figure in figures.values():
            figure['root'] = OUT
        suffix = '_Anonymous' if args.anonymous else ''
        folder = '匿名版' if args.anonymous else '正式版'
        target = args.output_root / folder / (name+suffix+'.docx')
        target.parent.mkdir(parents=True, exist_ok=True)
        previous = OUT / folder / target.name
        # Capture an existing file before the default in-place conversion replaces it.
        equation_source = previous
        if target.resolve() == previous.resolve() and previous.exists():
            import shutil
            equation_source = ROOT / 'Temp/PaperEquationSources' / folder / target.name
            equation_source.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(previous, equation_source)
        pypandoc.convert_text(content, 'docx', format='markdown+tex_math_dollars', outputfile=str(target))
        style_document(target, english, figures, AUTHORS, anonymous=args.anonymous)
        preserve_display_equations(target, equation_source)
        checklist.extend(['\n## '+name+'\n', notes])
        d = Document(target)
        print(f'{target.name}: paragraphs={len(d.paragraphs)}, tables={len(d.tables)}, display_equations={len(d.element.xpath(".//m:oMathPara"))}')
    (OUT / 'AUTHOR_CHECKLIST.md').write_text('\n\n'.join(checklist), encoding='utf-8')


if __name__ == '__main__':
    main()
