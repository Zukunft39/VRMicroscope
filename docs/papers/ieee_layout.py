"""Editable IEEE-style Word layout, including grouped figures and author metadata."""
from copy import deepcopy
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

WIDTH = 7.25
COLUMN = 3.5
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII']


def font(run, size=10, bold=None, italic=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph(p, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0,
              before=0, after=0, keep=False):
    p.alignment = align
    f = p.paragraph_format
    f.first_line_indent = Inches(indent)
    f.left_indent = f.right_indent = Inches(0)
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = 1
    f.keep_with_next = keep
    f.widow_control = True
    for r in p.runs:
        font(r, size)


def section_end(element, base, columns):
    """The section properties describe the preceding content, as Word requires."""
    p = OxmlElement('w:p')
    pr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    for key in ['before', 'after']:
        spacing.set(qn('w:' + key), '0')
    spacing.set(qn('w:line'), '20')
    spacing.set(qn('w:lineRule'), 'exact')
    pr.append(spacing)
    sect = deepcopy(base)
    kind = sect.find(qn('w:type'))
    if kind is None:
        kind = OxmlElement('w:type')
        sect.insert(0, kind)
    kind.set(qn('w:val'), 'continuous')
    cols = sect.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sect.append(cols)
    cols.set(qn('w:num'), str(columns))
    cols.set(qn('w:space'), '360')
    pr.append(sect)
    p.append(pr)
    element.addnext(p)


def borderless(table, width):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        item = OxmlElement('w:' + edge)
        item.set(qn('w:val'), 'nil')
        borders.append(item)
    props.append(borders)
    props.find(qn('w:tblW')).set(qn('w:w'), str(int(width * 1440)))
    props.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
    for col in table.columns:
        col.width = Inches(width / len(table.columns))
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.width = Inches(width / len(table.columns))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def authors_block(doc, anchor, metadata, anonymous, english=True):
    anchor.clear()
    paragraph(anchor, 11, WD_ALIGN_PARAGRAPH.CENTER, after=5, keep=True)
    if anonymous:
        font(anchor.add_run('Anonymous authors' if english else '匿名作者'), 11)
        return anchor._p
    people = metadata['authors']
    for i, person in enumerate(people):
        if i and i != (len(people) + 1) // 2:
            anchor.add_run(', ')
        # Break long author lists into balanced rows without changing order.
        if i == (len(people) + 1) // 2:
            anchor.add_run().add_break()
        font(anchor.add_run(person['name']), 11)
        corr_star = '*' if person.get('corresponding') else ''
        mark = anchor.add_run(f"{person['affiliation']}{corr_star}")
        font(mark, 8)
        mark.font.superscript = True
    table = doc.add_table(rows=1, cols=len(metadata['affiliations']))
    anchor._p.addnext(table._tbl)
    borderless(table, WIDTH)
    for cell, affiliation in zip(table.rows[0].cells, metadata['affiliations']):
        p = cell.paragraphs[0]
        paragraph(p, 9, WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        mark = p.add_run(str(affiliation['id']))
        font(mark, 7)
        mark.font.superscript = True
        font(p.add_run(affiliation['institution']), 9, italic=True)
        p = cell.add_paragraph(affiliation['country'])
        paragraph(p, 9, WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
        for person in people:
            if person['affiliation'] == affiliation['id']:
                p = cell.add_paragraph(person['email'])
                paragraph(p, 8, WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        cell.paragraphs[-1].paragraph_format.space_after = Pt(6)
    corr_authors = [p for p in people if p.get('corresponding')]
    if corr_authors:
        corr_p = doc.add_paragraph()
        table._tbl.addnext(corr_p._p)
        paragraph(corr_p, 8, WD_ALIGN_PARAGRAPH.CENTER, after=6, keep=True)
        if english:
            label = '*Corresponding authors: ' if len(corr_authors) > 1 else '*Corresponding author: '
        else:
            label = '*通讯作者：'
        names_and_emails = '; '.join(f"{p['name']} ({p['email']})" for p in corr_authors)
        run = corr_p.add_run(label + names_and_emails)
        font(run, 8, italic=True)
        return corr_p._p
    return table._tbl


def figure_block(doc, anchor, definition, base):
    panels = definition['panels']
    if len(panels) == 1:
        source = definition['root'] / panels[0][0]
        image_p = anchor.insert_paragraph_before()
        paragraph(image_p, 8, WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        picture = image_p.add_run().add_picture(str(source), width=Inches(7.1))
        svg = source.with_suffix('.svg')
        if svg.exists():
            part = Part(PackURI('/word/media/fig03_architecture.svg'),
                        'image/svg+xml', svg.read_bytes(), doc.part.package)
            rid = doc.part.relate_to(part, RT.IMAGE)
            blip = picture._inline.xpath('.//a:blip')[0]
            extensions = OxmlElement('a:extLst')
            extension = OxmlElement('a:ext')
            extension.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
            ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
            vector = etree.Element('{' + ns + '}svgBlip', nsmap={'asvg': ns})
            vector.set(qn('r:embed'), rid)
            extension.append(vector)
            extensions.append(extension)
            blip.append(extensions)
        anchor.text = definition['caption']
        paragraph(anchor, 8, before=3, after=8)
        anchor.paragraph_format.keep_together = True
        section_end(image_p._p.getprevious(), base, 2)
        section_end(anchor._p, base, 1)
        return
    optical = len(panels) == 5
    table = doc.add_table(rows=2 if optical else (len(panels) + 1) // 2,
                          cols=6 if optical else 2)
    anchor._p.addprevious(table._tbl)
    borderless(table, WIDTH)
    for idx, (source, label) in enumerate(panels):
        if optical:
            if idx < 3:
                cell = table.cell(0, idx*2).merge(table.cell(0, idx*2+1))
                image_width = 2.23
            else:
                col = (idx-3)*3
                cell = table.cell(1, col).merge(table.cell(1, col+2))
                image_width = 3.43
        else:
            cell = table.cell(idx // 2, idx % 2)
            image_width = 3.43
        p = cell.paragraphs[0]
        paragraph(p, 8, WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        p.add_run().add_picture(str(definition['root'] / source), width=Inches(image_width))
        p = cell.add_paragraph(label)
        paragraph(p, 8, WD_ALIGN_PARAGRAPH.CENTER, after=5, keep=True)
    # An odd last panel is centered across both columns at one-column scale.
    if len(panels) % 2 and not optical:
        cell = table.cell(len(panels) // 2, 0).merge(table.cell(len(panels) // 2, 1))
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    anchor.text = definition['caption']
    paragraph(anchor, 8, before=3, after=8)
    anchor.paragraph_format.keep_together = True
    previous = table._tbl.getprevious()
    section_end(previous, base, 2)
    section_end(anchor._p, base, 1)


def style_document(path, english, figures, metadata, anonymous=False):
    doc = Document(path)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.75), Inches(1)
    sec.left_margin = sec.right_margin = Inches(.625)
    sec.header_distance = sec.footer_distance = Inches(.3)
    for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
        if name not in doc.styles:
            continue
        s = doc.styles[name]
        s.font.name, s.font.size = 'Times New Roman', Pt(10)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
        s.paragraph_format.space_before = s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.line_spacing = 1
        s.paragraph_format.widow_control = True
    for name in ['Heading 1', 'Heading 2', 'Heading 3', 'Title']:
        s = doc.styles[name]
        s.font.name, s.font.size = 'Times New Roman', Pt(10)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.keep_with_next = name.startswith('Heading')
        s.paragraph_format.page_break_before = False
    # Avoid colored hyperlinks in a printed proceedings manuscript.
    if 'Hyperlink' in doc.styles:
        doc.styles['Hyperlink'].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles['Hyperlink'].font.underline = False
    cols = sec._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sec._sectPr.append(cols)
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:space'), '360')
    kind = sec._sectPr.find(qn('w:type'))
    if kind is None:
        kind = OxmlElement('w:type')
        sec._sectPr.insert(0, kind)
    kind.set(qn('w:val'), 'continuous')
    base = deepcopy(sec._sectPr)
    paragraphs = list(doc.paragraphs)
    tables = list(doc.tables)
    title = paragraphs[0]
    paragraph(title, 24 if english else 22, WD_ALIGN_PARAGRAPH.CENTER, after=12, keep=True)
    for r in title.runs:
        r.bold = False
    last_author = authors_block(doc, paragraphs[1], metadata, anonymous, english=english)
    section_end(last_author, base, 1)
    for p in paragraphs[2:]:
        t = p.text.strip()
        if re.fullmatch(r'FIGUREGROUP\d+', t):
            continue
        if p.style.name == 'Heading 2':
            match = re.match(r'(\d+)\.\s*(.*)', t)
            if match:
                t = ROMAN[int(match.group(1)) - 1] + '. ' + match.group(2)
            p.text = t.upper() if english else t
            paragraph(p, 10, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=5, keep=True)
            for r in p.runs:
                r.bold = False
        elif p.style.name == 'Heading 3':
            match = re.match(r'\d+\.(\d+)\s*(.*)', t)
            if match:
                p.text = chr(64 + int(match.group(1))) + '. ' + match.group(2)
            paragraph(p, 10, WD_ALIGN_PARAGRAPH.LEFT, before=7, after=3, keep=True)
            for r in p.runs:
                r.bold, r.italic = False, True
        elif t.startswith(('Abstract—', '摘要—', 'Index Terms—', '关键词—')):
            paragraph(p, 9, after=5)
            for r in p.runs:
                r.bold = True
            if t.startswith(('Index Terms—', '关键词—')):
                p.runs[0].italic = True
        elif re.match(r'^(?:Table|表)\s*\d+[:.：]', t):
            match = re.match(r'^(?:Table|表)\s*(\d+)[:.：]\s*(.*)', t)
            p.text = ('TABLE ' if english else '表 ') + ROMAN[int(match.group(1)) - 1] + '\n' + match.group(2)
            paragraph(p, 8, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
            for r in p.runs:
                r.bold = False
        elif t.startswith('FIGURE_SLOT_'):
            p.text = '[Architecture artwork pending]' if english else '[系统架构图待补]'
            paragraph(p, 9, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=5, keep=True)
        elif re.match(r'^(?:Fig\.|图)\s*\d+\.', t):
            paragraph(p, 8, before=3, after=8)
        elif re.match(r'^(?:\*?Note:|\*?注[：:])', t):
            paragraph(p, 8, WD_ALIGN_PARAGRAPH.JUSTIFY, before=2, after=6)
            for r in p.runs:
                r.italic = True
        elif re.match(r'^\[\d+\]', t):
            paragraph(p, 8, WD_ALIGN_PARAGRAPH.LEFT, after=3)
            p.paragraph_format.left_indent = Inches(.2)
            p.paragraph_format.first_line_indent = Inches(-.2)
        elif p._p.xpath('.//m:oMathPara'):
            paragraph(p, 9, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)
            for mr in p._p.xpath('.//m:r'):
                pr = mr.find(qn('w:rPr'))
                if pr is None:
                    pr = OxmlElement('w:rPr')
                    mr.append(pr)
                sz = pr.find(qn('w:sz'))
                if sz is None:
                    sz = OxmlElement('w:sz')
                    pr.append(sz)
                sz.set(qn('w:val'), '18')
        else:
            paragraph(p, 10, indent=.14)
            if p._p.xpath('./w:pPr/w:numPr'):
                p.paragraph_format.first_line_indent = Inches(-.12)
                p.paragraph_format.left_indent = Inches(.24)
    # Data tables: booktabs-style rules, repeating headers, no vertical grid.
    for table in tables:
        wide = len(table.columns) >= 3
        borderless(table, WIDTH if wide else COLUMN)
        borders = table._tbl.tblPr.find(qn('w:tblBorders'))
        for edge in ['top', 'bottom']:
            item = borders.find(qn('w:' + edge))
            item.set(qn('w:val'), 'single')
            item.set(qn('w:sz'), '6')
        for ridx, row in enumerate(table.rows):
            if ridx == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    paragraph(p, 8, WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2, keep=ridx<len(table.rows)-1)
                    for r in p.runs:
                        r.bold = ridx == 0
                if ridx == 0:
                    b = OxmlElement('w:tcBorders')
                    rule = OxmlElement('w:bottom')
                    rule.set(qn('w:val'), 'single')
                    rule.set(qn('w:sz'), '4')
                    b.append(rule)
                    cell._tc.get_or_add_tcPr().append(b)
        if wide:
            caption = table._tbl.getprevious()
            section_end(caption.getprevious(), base, 2)
            section_end(table._tbl, base, 1)
    for p in paragraphs:
        match = re.fullmatch(r'FIGUREGROUP(\d+)', p.text.strip())
        if match:
            figure_block(doc, p, figures[int(match.group(1))], base)
    # Keep table numbering in prose consistent with Roman table captions.
    for p in doc.paragraphs:
        for r in p.runs:
            updated = re.sub(r'\bTable ([1-4])\b', lambda m: 'Table ' + ROMAN[int(m.group(1)) - 1], r.text)
            updated = re.sub(r'表 ([1-4])(?=[，、。 中及和之]|$)', lambda m: '表 ' + ROMAN[int(m.group(1)) - 1], updated)
            if updated != r.text:
                r.text = updated
    doc.core_properties.title = title.text
    doc.core_properties.author = 'Anonymous' if anonymous else '; '.join(a['name'] for a in metadata['authors'])
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    doc.save(path)
